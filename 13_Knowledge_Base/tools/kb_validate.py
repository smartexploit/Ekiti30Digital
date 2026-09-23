#!/usr/bin/env python3
"""
kb_validate.py: validate Ask Ekiti knowledge-base documents and build kb_manifest.csv

Implements the rules in Ask_Ekiti_KB_Spec.md (A2-A5, A9) against the topic
folders 01_History ... 09_Statistics and the Source Inventory workbook.

Only the Python standard library is needed. Reading the .xlsx registry needs
openpyxl (pip install openpyxl); otherwise export the "Source Inventory" sheet
to CSV and pass it with --registry.

Usage (from the repository root):
    python 13_Knowledge_Base/tools/kb_validate.py
    python 13_Knowledge_Base/tools/kb_validate.py --strict
    python 13_Knowledge_Base/tools/kb_validate.py --selftest

Exit code: 0 = no errors, 1 = errors found (or warnings with --strict).

Front matter format supported: simple "key: value" lines, inline lists like
[SRC-001, SRC-002], and "# comments". Multi-line or nested YAML is not supported.
"""
import argparse
import csv
import datetime
import hashlib
import os
import re
import sys
import tempfile

TOPIC_DIRS = ["01_History", "02_LGAs", "03_Timeline", "04_Tourism", "05_Culture",
              "06_Education", "07_Health", "08_Agriculture", "09_Statistics"]
CATEGORY_FOLDERS = {
    "history": ["01_History", "03_Timeline"],
    "government": ["01_History/government"],
    "lgas": ["02_LGAs"],
    "tourism": ["04_Tourism"],
    "culture": ["05_Culture"],
    "education": ["06_Education"],
    "agriculture": ["08_Agriculture"],
    "health": ["07_Health"],
    "statistics": ["09_Statistics"],
}
DOC_TYPES = {"entity", "event", "dataset", "background"}
STATUSES = {"draft", "needs_review", "conflict", "verified", "retired"}
TIERS = {"A", "B", "C", "D"}
LANGUAGES = {"en", "yo"}
REQUIRED = ["id", "title", "category", "doc_type", "source_ids", "source_name", "source_url",
            "publication_date", "source_tier", "status", "period_covered", "language"]
SKIP_NAMES = {"README.md", "CONTRIBUTING.md", "CODE_OF_CONDUCT.md"}
MANIFEST_COLUMNS = ["id", "path", "category", "doc_type", "status", "source_tier", "source_ids",
                    "publication_date", "last_verified", "verified_by", "language", "translation_of",
                    "period_covered", "file_sha256", "ingestible", "reason"]


# ---------------------------------------------------------------- parsing
def strip_comment(val):
    val = val.strip()
    if val[:1] in ("'", '"'):
        q = val[0]
        end = val.find(q, 1)
        if end != -1:
            return val[:end + 1]
    return re.sub(r"\s+#.*$", "", val)


def parse_value(val):
    val = strip_comment(val).strip()
    if val == "":
        return ""
    if val.startswith("[") and val.endswith("]"):
        inner = val[1:-1].strip()
        if not inner:
            return []
        return [x.strip().strip("'\"") for x in inner.split(",") if x.strip()]
    if val.lower() in ("true", "false"):
        return val.lower() == "true"
    if val[:1] in ("'", '"') and val[-1:] == val[:1] and len(val) >= 2:
        return val[1:-1]
    return val


def parse_front_matter(text):
    """Return (dict or None, body, error message or None)."""
    lines = text.splitlines()
    if not lines or lines[0].strip() != "---":
        return None, text, "no front matter"
    end = None
    for i in range(1, len(lines)):
        if lines[i].strip() == "---":
            end = i
            break
    if end is None:
        return None, text, "front matter is not closed with ---"
    fm = {}
    for ln in lines[1:end]:
        if not ln.strip() or ln.lstrip().startswith("#"):
            continue
        if ln[0] in (" ", "\t"):
            return None, text, f"indented or multi-line value not supported: {ln.strip()!r}"
        if ":" not in ln:
            return None, text, f"cannot parse front-matter line: {ln!r}"
        key, val = ln.split(":", 1)
        fm[key.strip()] = parse_value(val)
    return fm, "\n".join(lines[end + 1:]), None


# ---------------------------------------------------------------- dates
PARTIAL_DATE = re.compile(r"^\d{4}(-\d{2}(-\d{2})?)?$")


def is_placeholder(s):
    s = str(s).strip()
    return s == "" or ("YYYY" in s.upper() and not re.search(r"\d{4}", s))


def parse_partial_date(s):
    """Return the earliest date the (possibly partial) value can mean, or None if invalid."""
    s = str(s).strip()
    if not PARTIAL_DATE.match(s):
        return None
    parts = [int(p) for p in s.split("-")]
    while len(parts) < 3:
        parts.append(1)
    try:
        return datetime.date(*parts)
    except ValueError:
        return None


def parse_full_date(s):
    s = str(s).strip()
    if not re.match(r"^\d{4}-\d{2}-\d{2}$", s):
        return None
    try:
        return datetime.date.fromisoformat(s)
    except ValueError:
        return None


# ---------------------------------------------------------------- registry
def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def load_registry(path):
    rows = []
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.DictReader(f))
    else:
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("openpyxl is needed to read the .xlsx registry (pip install openpyxl), "
                     "or export the 'Source Inventory' sheet to CSV and use --registry file.csv")
        wb = load_workbook(path, data_only=True)
        ws = wb["Source Inventory"]
        header = [c.value for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0]:
                rows.append({h: ("" if v is None else v) for h, v in zip(header, r)})
    reg = {}
    for r in rows:
        sid = norm(r.get("Source ID"))
        if not sid:
            continue
        reg[sid] = {
            "stage": norm(r.get("Record stage")),
            "status": norm(r.get("Status")),
            "tier": norm(r.get("Tier")),
            "title": norm(r.get("Source title")),
            "location": norm(r.get("URL or document location")),
        }
    return reg


# ---------------------------------------------------------------- checks
def validate_doc(doc, registry, today):
    fm, path = doc["fm"], doc["path"]
    errors, warnings = [], []
    E, W = errors.append, warnings.append

    missing = []
    for k in REQUIRED:
        if k == "source_ids":
            if not isinstance(fm.get(k), list) or not fm.get(k):
                missing.append(k)
        elif fm.get(k) in (None, ""):
            missing.append(k)
    for k in missing:
        E(f"missing required field: {k}")
    if missing and ("status" in missing or "id" in missing):
        return errors, warnings

    status = fm.get("status", "")
    category = fm.get("category", "")
    doc_type = fm.get("doc_type", "")
    tier = fm.get("source_tier", "")
    lang = fm.get("language", "")
    period = str(fm.get("period_covered", "")).strip()
    sids = fm.get("source_ids") if isinstance(fm.get("source_ids"), list) else []

    if not re.fullmatch(r"[a-z0-9]+(-[a-z0-9]+)*", str(fm.get("id", ""))):
        E("id must be lowercase letters, digits and hyphens (e.g. lga-ado-ekiti)")
    if category and category not in CATEGORY_FOLDERS:
        E(f"category '{category}' is not one of: {', '.join(sorted(CATEGORY_FOLDERS))}")
    if doc_type and doc_type not in DOC_TYPES:
        E(f"doc_type '{doc_type}' is not one of: {', '.join(sorted(DOC_TYPES))}")
    if status and status not in STATUSES:
        E(f"status '{status}' is not one of: {', '.join(sorted(STATUSES))}")
    if tier and tier not in TIERS:
        E(f"source_tier '{tier}' is not one of A, B, C, D")
    if lang and lang not in LANGUAGES:
        E(f"language '{lang}' is not one of: en, yo")

    # ---- folder rules (A2)
    prefixes = CATEGORY_FOLDERS.get(category, [])
    if prefixes and not any(path == p or path.startswith(p + "/") for p in prefixes):
        E(f"category '{category}' files belong under {' or '.join(prefixes)}")
    if path.startswith("01_History/government/") and category != "government":
        E("files in 01_History/government/ must have category: government")
    in_pre = path.startswith("01_History/pre-1996/")
    if period.lower() == "pre-1996":
        if doc_type != "background":
            E("pre-1996 documents must have doc_type: background")
        if not in_pre:
            E("pre-1996 documents must live in 01_History/pre-1996/")
    elif in_pre:
        E("files in 01_History/pre-1996/ must have period_covered: pre-1996")
    if path.startswith("03_Timeline/"):
        if doc_type != "event":
            E("03_Timeline is for events only (doc_type: event)")
        if period.lower() == "pre-1996":
            E("pre-1996 content must stay out of the 1996-2026 timeline")
    if period.lower() != "pre-1996" and period:
        m = re.fullmatch(r"(\d{4})(?:-(\d{4}))?", period)
        if not m:
            E("period_covered must be a year (1996), a range (1996-2026) or pre-1996")
        else:
            y1 = int(m.group(1))
            y2 = int(m.group(2) or m.group(1))
            if y1 < 1996 or y2 < 1996:
                E("years before 1996 must use period_covered: pre-1996")
            if y1 > y2:
                E("period_covered range is reversed")
            if y2 > today.year:
                W("period_covered ends after the current year")

    # ---- sources (A4, A5, A6)
    first = None
    for sid in sids:
        if not re.fullmatch(r"SRC-\d{3}", sid):
            E(f"source id '{sid}' is not in the form SRC-###")
            continue
        r = registry.get(sid)
        if r is None:
            E(f"{sid} is not in the Source Inventory")
            continue
        first = first or (sid, r)
        if r["stage"] == "Source needed":
            (E if status == "verified" else W)(f"{sid} is still 'Source needed' in the inventory")
        if r["status"] == "Conflict":
            (E if status == "verified" else W)(f"{sid} is in Conflict; resolve it before verifying (spec A6)")
        if r["tier"] == "D":
            (E if status in ("verified", "needs_review") else W)(f"{sid} is Tier D, which is not used at launch")
        if status == "verified" and r["status"] not in ("Verified",):
            W(f"{sid} record status is '{r['status']}', not 'Verified'")
    if first:
        sid, r = first
        if r["stage"] == "Identified":
            if norm(fm.get("source_name")) != r["title"]:
                E(f"source_name does not match the inventory title for {sid}")
            if norm(fm.get("source_url")) != r["location"]:
                E(f"source_url does not match the inventory location for {sid}")
        if tier and r["tier"] and tier != r["tier"]:
            W(f"source_tier {tier} differs from {sid} tier {r['tier']} in the inventory")
    if "wikipedia" in str(fm.get("source_url", "")).lower():
        E("Wikipedia is not an accepted source")

    # ---- dates
    pub = str(fm.get("publication_date", "")).strip()
    if pub and pub.lower() != "not stated":
        d = parse_partial_date(pub)
        if is_placeholder(pub) or d is None:
            E("publication_date must be YYYY, YYYY-MM, YYYY-MM-DD or 'Not stated'")
        elif d > today:
            E("publication_date is in the future")
    lv = str(fm.get("last_verified", "")).strip()
    vb = str(fm.get("verified_by", "")).strip()
    note = str(fm.get("verification_note", "")).strip()
    exc = fm.get("tier_c_exception") is True
    if status == "verified":
        if not vb:
            E("verified documents need verified_by")
        d = parse_full_date(lv)
        if is_placeholder(lv) or d is None:
            E("verified documents need a real last_verified date (YYYY-MM-DD), not the placeholder")
        elif d > today:
            E("last_verified is in the future")
        if tier == "C" and len(sids) < 2 and not (exc and note):
            E("Tier C needs two independent sources, or tier_c_exception: true with a verification_note (A7)")
        if tier == "D":
            E("Tier D is not used at launch")
    else:
        if lv and not is_placeholder(lv):
            W("last_verified is set but status is not verified")
        if vb:
            W("verified_by is set but status is not verified")
    if exc and not note:
        (E if status == "verified" else W)("tier_c_exception needs a verification_note")

    # ---- language (A9 rule 9)
    if lang == "yo":
        if not fm.get("translation_of"):
            E("Yoruba documents need translation_of (id of the English original)")
        if status == "verified" and not fm.get("translation_reviewed_by"):
            E("verified Yoruba documents need translation_reviewed_by (Yoruba reviewer)")
    elif lang == "en" and fm.get("translation_of"):
        W("translation_of is set on an English document")

    # ---- body
    body = doc["body"]
    if not re.search(r"^##\s+Summary", body, re.M):
        W("body has no '## Summary' section")
    if not re.search(r"^##\s+Sources", body, re.M):
        W("body has no '## Sources' section")
    if re.search(r"\bTODO\b|\bTBD\b", body):
        W("body contains TODO/TBD")
    return errors, warnings


# ---------------------------------------------------------------- run
def collect_docs(root):
    docs, unparsed = [], []
    for top in TOPIC_DIRS:
        base = os.path.join(root, top)
        if not os.path.isdir(base):
            continue
        for dp, dn, fn in os.walk(base):
            dn.sort()
            for name in sorted(fn):
                if not name.endswith(".md") or name in SKIP_NAMES or name.startswith("_"):
                    continue
                full = os.path.join(dp, name)
                rel = os.path.relpath(full, root).replace(os.sep, "/")
                raw = open(full, "rb").read()
                text = raw.decode("utf-8", errors="replace")
                fm, body, err = parse_front_matter(text)
                if fm is None:
                    unparsed.append((rel, err))
                else:
                    docs.append({"path": rel, "fm": fm, "body": body,
                                 "sha": hashlib.sha256(raw).hexdigest()})
    return docs, unparsed


def run(root, registry_path, manifest_path, strict, write_manifest, today, quiet=False):
    registry = load_registry(registry_path)
    docs, unparsed = collect_docs(root)
    results, ids = {}, {}
    for d in docs:
        d["errors"], d["warnings"] = validate_doc(d, registry, today)
        i = d["fm"].get("id")
        if i:
            ids.setdefault(i, []).append(d["path"])
    for i, paths in ids.items():
        if len(paths) > 1:
            for d in docs:
                if d["fm"].get("id") == i:
                    d["errors"].append(f"duplicate id '{i}' (also in {', '.join(p for p in paths if p != d['path'])})")
    known = set(ids)
    titles = {}
    for d in docs:
        if d["fm"].get("language") == "yo" and d["fm"].get("translation_of") and d["fm"]["translation_of"] not in known:
            d["errors"].append(f"translation_of '{d['fm']['translation_of']}' does not match any document id")
        key = (d["fm"].get("category"), norm(d["fm"].get("title")).lower(), d["fm"].get("language"))
        titles.setdefault(key, []).append(d["path"])
    for key, paths in titles.items():
        if key[1] and len(paths) > 1:
            for d in docs:
                if d["path"] in paths:
                    d["warnings"].append("another document has the same title (possible duplicate)")

    n_err = n_warn = 0
    lines = []
    for rel, err in unparsed:
        lines.append(f"WARNING {rel}: {err} (not treated as a KB document)")
        n_warn += 1
    for d in docs:
        for m in d["errors"]:
            lines.append(f"ERROR   {d['path']}: {m}")
            n_err += 1
        for m in d["warnings"]:
            lines.append(f"WARNING {d['path']}: {m}")
            n_warn += 1

    rows = []
    for d in sorted(docs, key=lambda x: x["path"]):
        fm = d["fm"]
        ok = fm.get("status") == "verified" and not d["errors"]
        reason = "" if ok else ("has errors" if d["errors"] else f"status is {fm.get('status', '?')}")
        rows.append({
            "id": fm.get("id", ""), "path": d["path"], "category": fm.get("category", ""),
            "doc_type": fm.get("doc_type", ""), "status": fm.get("status", ""),
            "source_tier": fm.get("source_tier", ""), "source_ids": ";".join(fm.get("source_ids") or []),
            "publication_date": fm.get("publication_date", ""), "last_verified": fm.get("last_verified", ""),
            "verified_by": fm.get("verified_by", ""), "language": fm.get("language", ""),
            "translation_of": fm.get("translation_of", ""), "period_covered": fm.get("period_covered", ""),
            "file_sha256": d["sha"], "ingestible": "yes" if ok else "no", "reason": reason,
        })
    if write_manifest:
        os.makedirs(os.path.dirname(os.path.abspath(manifest_path)), exist_ok=True)
        with open(manifest_path, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=MANIFEST_COLUMNS)
            w.writeheader()
            w.writerows(rows)
    if not quiet:
        for ln in lines:
            print(ln)
        n_ing = sum(1 for r in rows if r["ingestible"] == "yes")
        n_ver = sum(1 for r in rows if r["status"] == "verified")
        print(f"\nDocuments: {len(docs)} | verified: {n_ver} | ingestible: {n_ing} | "
              f"errors: {n_err} | warnings: {n_warn}")
        if write_manifest:
            print(f"Manifest written: {manifest_path}")
    return docs, rows, n_err, n_warn


# ---------------------------------------------------------------- selftest
def selftest():
    today = datetime.date(2026, 9, 21)
    tmp = tempfile.mkdtemp()
    reg = os.path.join(tmp, "reg.csv")
    with open(reg, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Source ID", "Record stage", "Source title", "URL or document location", "Tier", "Status"])
        w.writerow(["SRC-001", "Identified", "Good Source", "https://example.org/a", "A", "Verified"])
        w.writerow(["SRC-002", "Source needed", "", "", "A", "Not started"])
        w.writerow(["SRC-008", "Identified", "Census", "https://example.org/c", "A", "Conflict"])
        w.writerow(["SRC-026", "Source needed", "", "", "D", "Blocked"])
        w.writerow(["SRC-030", "Identified", "News item", "https://example.org/n", "C", "Verified"])

    def doc(path, **over):
        fm = dict(id="doc-" + os.path.basename(path)[:-3].lower(), title="T " + path, category="history",
                  doc_type="event", source_ids="[SRC-001]", source_name="Good Source",
                  source_url="https://example.org/a", publication_date="2021-10-01", source_tier="A",
                  status="verified", verified_by="Reviewer", last_verified="2026-09-20",
                  period_covered="1996", language="en")
        fm.update(over)
        lines = ["---"] + [f"{k}: {v}" for k, v in fm.items() if v is not None] + ["---",
                 "## Summary", "x", "## Sources", "- [S1] x"]
        full = os.path.join(tmp, path)
        os.makedirs(os.path.dirname(full), exist_ok=True)
        open(full, "w", encoding="utf-8").write("\n".join(lines))

    doc("01_History/good.md")
    doc("01_History/future.md", last_verified="2026-09-25")
    doc("01_History/placeholder.md", last_verified="YYYY-MM-DD")
    doc("01_History/conflict.md", source_ids="[SRC-008]", source_name="Census", source_url="https://example.org/c")
    doc("03_Timeline/pre.md", period_covered="pre-1996", doc_type="background")
    doc("01_History/gov.md", category="government")
    doc("01_History/tierc.md", source_tier="C", source_ids="[SRC-030]", source_name="News item",
        source_url="https://example.org/n")
    doc("01_History/yo.md", language="yo")
    doc("01_History/unknown.md", source_ids="[SRC-999]")
    doc("01_History/missing.md", source_name=None)
    doc("01_History/mismatch.md", source_name="Other Title")
    doc("01_History/tierd.md", source_ids="[SRC-026]", source_tier="D", source_name="", source_url="")
    doc("01_History/pre-1996/ok.md", period_covered="pre-1996", doc_type="background", status="draft",
        verified_by="", last_verified="")
    doc("01_History/pre-1996/bad.md", period_covered="1996", doc_type="background", status="draft",
        verified_by="", last_verified="")
    docs, rows, n_err, _ = run(tmp, reg, os.path.join(tmp, "m.csv"), False, True, today, quiet=True)
    by = {d["path"]: d for d in docs}
    expect = {
        "01_History/good.md": None,
        "01_History/future.md": "in the future",
        "01_History/placeholder.md": "real last_verified",
        "01_History/conflict.md": "Conflict",
        "03_Timeline/pre.md": "pre-1996",
        "01_History/gov.md": "belong under",
        "01_History/tierc.md": "Tier C needs two",
        "01_History/yo.md": "translation_of",
        "01_History/unknown.md": "not in the Source Inventory",
        "01_History/missing.md": "missing required field",
        "01_History/mismatch.md": "does not match",
        "01_History/tierd.md": "Tier D",
        "01_History/pre-1996/ok.md": None,
        "01_History/pre-1996/bad.md": "must have period_covered: pre-1996",
    }
    failures = []
    for p, needle in expect.items():
        errs = " | ".join(by[p]["errors"])
        if needle is None and errs:
            failures.append(f"{p}: expected no errors, got: {errs}")
        if needle is not None and needle not in errs:
            failures.append(f"{p}: expected error containing '{needle}', got: {errs or 'none'}")
    ing = {r["path"]: r["ingestible"] for r in rows}
    if ing.get("01_History/good.md") != "yes":
        failures.append("good.md should be ingestible")
    if any(v == "yes" for p, v in ing.items() if p != "01_History/good.md"):
        failures.append("only good.md should be ingestible")
    if failures:
        print("SELFTEST FAILED")
        for f in failures:
            print("  -", f)
        return 1
    print(f"SELFTEST PASSED ({len(expect)} cases)")
    return 0


def main():
    ap = argparse.ArgumentParser(description="Validate Ask Ekiti KB documents and build the manifest.")
    ap.add_argument("--root", default=".", help="repository root (default: current folder)")
    ap.add_argument("--registry", default=None,
                    help="Source Inventory .xlsx or .csv (default: 14_Source_Documents/Ask_Ekiti_Source_Inventory.xlsx)")
    ap.add_argument("--manifest", default=None, help="output CSV (default: 13_Knowledge_Base/kb_manifest.csv)")
    ap.add_argument("--no-manifest", action="store_true", help="validate only; do not write the manifest")
    ap.add_argument("--strict", action="store_true", help="treat warnings as errors")
    ap.add_argument("--today", default=None, help="override today's date (YYYY-MM-DD), for testing")
    ap.add_argument("--selftest", action="store_true", help="run the built-in test cases")
    a = ap.parse_args()
    if a.selftest:
        sys.exit(selftest())
    today = datetime.date.fromisoformat(a.today) if a.today else datetime.date.today()
    registry = a.registry or os.path.join(a.root, "14_Source_Documents", "Ask_Ekiti_Source_Inventory.xlsx")
    manifest = a.manifest or os.path.join(a.root, "13_Knowledge_Base", "kb_manifest.csv")
    if not os.path.exists(registry):
        sys.exit(f"registry not found: {registry}")
    _, _, n_err, n_warn = run(a.root, registry, manifest, a.strict, not a.no_manifest, today)
    sys.exit(1 if n_err or (a.strict and n_warn) else 0)


if __name__ == "__main__":
    main()
